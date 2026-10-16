"""Document loading functions for supported file types (PDF, DOCX, TXT, Markdown).

Each loader returns a list of LangChain Document objects with normalized
metadata: source, file_name, file_type (and page, for PDFs).
"""

from pathlib import Path
from typing import Dict, List

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from ingestion.table_repair import repair_wide_table_columns
from utils.logging_config import get_logger

logger = get_logger(__name__)

_UNREADABLE_MESSAGE = "Unable to read this document. Please check that the file is valid."
_EMPTY_MESSAGE = "The uploaded document is empty and cannot be processed."


class DocumentParsingError(Exception):
    """Raised when a document cannot be parsed into usable text."""


def _base_metadata(file_path: Path, file_type: str) -> Dict[str, str]:
    return {
        "source": str(file_path),
        "file_name": file_path.name,
        "file_type": file_type,
    }


def load_pdf_document(file_path: Path) -> List[Document]:
    """Load a PDF file, returning one Document per page with page numbers.

    Each page's extracted text is additionally passed through
    repair_wide_table_columns(), which only ever changes a page when a
    conservative, evidence-based signature (a wide table's last column
    disconnected from its own rows) is detected -- see
    ingestion/table_repair.py. Every other page, including every page of
    documents that never exhibit this defect, comes back byte-identical.
    """
    try:
        reader = PdfReader(str(file_path))
        page_texts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            try:
                text = repair_wide_table_columns(page, text)
            except Exception:
                logger.warning(
                    "Wide-table column repair failed for a page; using the "
                    "original extracted text for that page instead.",
                    exc_info=True,
                )
            page_texts.append(text)
    except Exception as exc:
        logger.exception("Failed to parse PDF file: %s", file_path)
        raise DocumentParsingError(_UNREADABLE_MESSAGE) from exc

    if not page_texts or not any(text.strip() for text in page_texts):
        raise DocumentParsingError(_EMPTY_MESSAGE)

    documents = []
    for index, text in enumerate(page_texts):
        metadata = _base_metadata(file_path, "pdf")
        metadata["page"] = index + 1
        documents.append(Document(page_content=text, metadata=metadata))
    return documents


def load_docx_document(file_path: Path) -> List[Document]:
    """Load a DOCX file, returning a single Document with the full text."""
    try:
        docx_file = DocxDocument(str(file_path))
        paragraphs = [p.text for p in docx_file.paragraphs if p.text.strip()]
    except Exception as exc:
        logger.exception("Failed to parse DOCX file: %s", file_path)
        raise DocumentParsingError(_UNREADABLE_MESSAGE) from exc

    text = "\n".join(paragraphs).strip()
    if not text:
        raise DocumentParsingError(_EMPTY_MESSAGE)

    return [Document(page_content=text, metadata=_base_metadata(file_path, "docx"))]


def load_txt_document(file_path: Path) -> List[Document]:
    """Load a plain-text file, returning a single Document."""
    try:
        text = file_path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        logger.exception("Failed to parse TXT file: %s", file_path)
        raise DocumentParsingError(_UNREADABLE_MESSAGE) from exc

    if not text.strip():
        raise DocumentParsingError(_EMPTY_MESSAGE)

    return [Document(page_content=text, metadata=_base_metadata(file_path, "txt"))]


def load_markdown_document(file_path: Path) -> List[Document]:
    """Load a Markdown file, returning a single Document with the raw content."""
    try:
        text = file_path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        logger.exception("Failed to parse Markdown file: %s", file_path)
        raise DocumentParsingError(_UNREADABLE_MESSAGE) from exc

    if not text.strip():
        raise DocumentParsingError(_EMPTY_MESSAGE)

    return [Document(page_content=text, metadata=_base_metadata(file_path, "md"))]


_LOADER_DISPATCH = {
    ".pdf": load_pdf_document,
    ".docx": load_docx_document,
    ".txt": load_txt_document,
    ".md": load_markdown_document,
}


def load_document(file_path: Path) -> List[Document]:
    """Load a document from disk, dispatching on its file extension."""
    loader = _LOADER_DISPATCH.get(file_path.suffix.lower())
    if loader is None:
        raise DocumentParsingError(
            "Unsupported file type. Please upload PDF, DOCX, TXT, or Markdown files."
        )
    return loader(file_path)
